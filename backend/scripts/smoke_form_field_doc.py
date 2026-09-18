"""真实合同模板冒烟测试（P1）：表单域 / #占位符 的端到端识别与回写。

用法（在 backend/ 下执行）：
    .venv/bin/python scripts/smoke_form_field_doc.py

它用 `data/hetong/2026-9-11 免烘烤捣打料合同模版2026.doc` 这份真实素材跑完整链路
（.doc 归一化 → 扫描 → 模拟 AI 输出 → 回写 → 自检 → docxtpl 渲染）。这是唯一能覆盖
「Word 表单域（FORMTEXT）+ 浮动文本框 + 合并单元格 + 中文重名标签」的用例，
合成 fixture 复现不出这些结构组合。素材不存在时脚本跳过（退出码 0），
这样它不会因为同事没同步 data/ 目录就把 CI 弄红。
"""

import asyncio
import sys
import zipfile
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import docxtpl  # noqa: E402

from app.services import docx_convert, template_rewrite, template_scan  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
SOURCE_DOC = ROOT / "data/hetong/2026-9-11 免烘烤捣打料合同模版2026.doc"
WORK = Path("var/form-field-smoke")

# 这份模板的已知变量：37 个表单域 + 2 个 #占位符
EXPECTED_FORM_FIELDS = 37
EXPECTED_HASH = 2
# 甲乙双方成对、原文标签完全相同的组（主体是唯一区分手段）
COLLIDING_PAIRS = [
    ("REGIST_ADDR", "SUPPLIER_ADDR"),
    ("BANK_NAME", "SUPPLIER_BANK_NAME"),
    ("BANK_ACCOUNT", "SUPPLIER_ACCOUNT_NUM"),
    ("SOCIAL_CREDIT_NUM", "SUPPLIER_SOCIAL_CREDIT_CODE"),
    ("DTK198", "DTK207"),
]

failures: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  ✅ {label}")
    else:
        failures.append(label)
        print(f"  ❌ {label} {detail}")


def all_texts(document) -> list[str]:
    """正文段落 + 表格单元格里的段落（表单域大量出现在表格中）。"""
    texts = [paragraph.text for paragraph in document.paragraphs]
    texts.extend(
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    return texts


def ai_like_fields(scan: template_scan.ScanResult) -> list[dict]:
    """模拟应用六的输出：具名槽位沿用模板原名做 key，主体取槽位提示。

    这里不做语义判断（那是模型的事），只关心"每个槽位一个字段、原名与落点都带回来"，
    这样才能验证回写链路。
    """
    fields = []
    for slot in scan.slots:
        if slot.kind not in ("form_field", "hash"):
            continue
        key = slot.name if template_scan.is_valid_key(slot.name) else f"hash_{slot.slot_id}"
        fields.append(
            {
                "key": key,
                "label": slot.label_hint or slot.name,
                "type": "text",
                "slot_id": slot.slot_id,
                "source_name": slot.name,
                "owner": slot.owner_hint,
            }
        )
    return fields


def main() -> int:
    if not SOURCE_DOC.is_file():
        print(f"⏭  跳过：真实素材不存在（{SOURCE_DOC}）")
        return 0

    WORK.mkdir(parents=True, exist_ok=True)

    print("1) .doc 归一化（LibreOffice）")
    source = asyncio.run(docx_convert.convert_legacy_doc(SOURCE_DOC, WORK))
    check("产出 .docx", source.is_file() and source.suffix == ".docx", str(source))

    print("\n2) 扫描")
    scan = template_scan.scan_document(source)
    kinds = Counter(item["kind"] for item in scan.candidates)
    print(f"   候选：{dict(kinds)}")
    form_slots = [slot for slot in scan.slots if slot.kind == "form_field"]
    hash_slots = [slot for slot in scan.slots if slot.kind == "hash"]
    check(f"识别出 {EXPECTED_FORM_FIELDS} 个表单域槽位", len(form_slots) == EXPECTED_FORM_FIELDS,
          f"实际 {len(form_slots)}")
    check(f"识别出 {EXPECTED_HASH} 个 #占位符槽位", len(hash_slots) == EXPECTED_HASH,
          f"实际 {len(hash_slots)}")
    # 合并单元格跨 4 列：早期实现会把同一个物理段落编成 4 个锚点，TAX_RATE 出 4 个候选
    tax_rate = [item for item in scan.candidates if item.get("name") == "TAX_RATE"]
    check("合并单元格里的 TAX_RATE 只出一个候选", len(tax_rate) == 1, f"实际 {len(tax_rate)}")
    # 浮动文本框：Choice 与 Fallback 两份副本，但大纲里只应出现一次（AI 不该看到两遍）
    bpo = [item for item in scan.candidates if item.get("name") == "BPO_ID"]
    check("文本框两份副本各出一条候选", len(bpo) == 2, f"实际 {len(bpo)}")
    check("回退分支不进大纲（AI 不会看到重复的空位）",
          sum(1 for item in bpo if item["shadow"]) == 1 and scan.outline.count("BPO_ID") == 1)
    # 同名变量合并：一个字段、多处落点
    check("同名变量已合并（PUR_BILL_TO_NAME / BPO_ID 各 2 处落点）",
          len(next(s for s in form_slots if s.name == "PUR_BILL_TO_NAME").candidate_ids) == 2
          and len(next(s for s in form_slots if s.name == "BPO_ID").candidate_ids) == 2)

    print("\n3) 主体消歧（中文标签重名）")
    owner = {slot.name: slot.owner_hint for slot in form_slots}
    for party_a, party_b in COLLIDING_PAIRS:
        check(f"{party_a} / {party_b} 主体可区分",
              owner.get(party_a) == "甲方" and owner.get(party_b) == "乙方",
              f"{party_a}={owner.get(party_a)} {party_b}={owner.get(party_b)}")

    print("\n4) 回写与自检")
    fields = ai_like_fields(scan)
    standardized = WORK / "standardized.docx"
    try:
        result = template_rewrite.rewrite_document(source, fields, standardized)
    except Exception as exc:  # noqa: BLE001 - 失败信息就是本次冒烟的结论
        check("回写与自检通过", False, getattr(exc, "message", str(exc)))
        print(f"\n❌ 失败 {len(failures)} 项：{failures}")
        return 1
    check("回写与自检通过", True)
    print(f"   占位符 {result['placeholder_count']} 处 / 字段 {len(result['fields'])} 个")
    check("无标签重名警告（主体已把重名区分开）",
          not [w for w in result["warnings"] if "重名" in w], str(result["warnings"]))

    labels = {field["key"]: field["label"] for field in result["fields"]}
    check("消歧后的标签带主体前缀",
          all(labels.get(a, "").startswith("甲方") and labels.get(b, "").startswith("乙方")
              for a, b in COLLIDING_PAIRS if a in labels and b in labels),
          str({k: v for k, v in labels.items() if k in dict(COLLIDING_PAIRS[:2])}))

    print("\n5) 产物校验")
    with zipfile.ZipFile(standardized) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    check("无 FORMTEXT 残留", "FORMTEXT" not in xml)
    check("无 fldChar 残留", "fldChar" not in xml)
    check("无 # 占位符残留", "#供应商名称#" not in xml and "#   年月日至年月日#" not in xml)
    check("文本框没被删掉（绘图元素仍在）", xml.count("txbxContent") == 4,
          f"txbxContent={xml.count('txbxContent')}")
    check("文本框两份副本都写了占位符", xml.count("{{ BPO_ID }}") == 2,
          f"实际 {xml.count('{{ BPO_ID }}')} 处")
    check("同名变量写全所有落点", xml.count("{{ PUR_BILL_TO_NAME }}") == 2,
          f"实际 {xml.count('{{ PUR_BILL_TO_NAME }}')} 处")

    texts = all_texts(Document(str(standardized)))
    check("中文标签保留（甲方：）",
          any(text.strip() == "甲方：{{ PUR_BILL_TO_NAME }}" for text in texts),
          str([t for t in texts if "PUR_BILL_TO_NAME" in t][:2]))
    check("表1 备注整句保留，只换掉税率",
          any(
              text == "本合同价格由不含税价格和{{ TAX_RATE }}%税组成，其中不含税价格为固定价格。"
              "若合同履行期间遇国家税率调整，价税合计相应调整。"
              for text in texts
          ),
          str([t for t in texts if "TAX_RATE" in t][:1]))
    check("备注行本身也还在（同格另一段）", any(text.strip() == "备注：" for text in texts))
    check("同段 9 个变量各自落位（生效期限）",
          any(text.count("{{ DTK") >= 9 for text in texts),
          str([t[:80] for t in texts if "DTK65" in t][:1]))
    check("#占位符变成了英文占位符", any("致{{" in text for text in texts),
          str([t for t in texts if "致" in t][:2]))

    print("\n6) docxtpl 渲染（证明产物真的能填）")
    document = docxtpl.DocxTemplate(str(standardized))
    document.render({field["key"]: f"值-{field['key']}" for field in result["fields"]})
    rendered = WORK / "rendered.docx"
    document.save(str(rendered))
    filled = "\n".join(all_texts(Document(str(rendered))))
    check("渲染结果已填入取值", "值-PUR_BILL_TO_NAME" in filled and "值-DTK9980" in filled)
    check("渲染结果不再有占位符", "{{" not in filled)

    print()
    if failures:
        print(f"❌ 失败 {len(failures)} 项：{failures}")
        return 1
    print("✅ 全部通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
