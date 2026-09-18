"""模板结构化冒烟测试（P1）：扫描 → 回写 → 自检 → 端到端转 PDF。

用法（在 backend/ 下执行）：
    .venv/bin/python scripts/smoke_structure.py
脚本会现场合成一份含四类"待填"痕迹的 .docx，不需要外部素材。
"""

import asyncio
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

from app.core.exceptions import BizError  # noqa: E402
from app.services import docx_convert, template_rewrite, template_scan, word_generator  # noqa: E402

WORK = Path("var/structure-smoke")

failures: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  ✅ {label}")
    else:
        failures.append(label)
        print(f"  ❌ {label} {detail}")


def build_fixture(path: Path) -> None:
    """合成模板：覆盖下划线填空 / 标签后为空 / 已有占位符 / 表格空单元格。"""
    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run("钢材采购合同模板")
    run.font.size = Pt(18)
    run.font.bold = True

    for label in ("甲方（采购方）：", "乙方（供应方）："):
        para = doc.add_paragraph()
        para.add_run(label)
        blank = para.add_run("            ")
        blank.font.underline = True  # 下划线填空：回写后占位符必须继承下划线

    doc.add_paragraph("合同金额：")

    para = doc.add_paragraph()
    para.add_run("签订日期：")
    para.add_run("{{ sign_date }}")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "钢材"
    doc.add_paragraph("备注：以上条款经双方确认后生效。")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def pick(candidates: list[dict], kind: str, index: int = 0) -> dict:
    matched = [item for item in candidates if item["kind"] == kind]
    if len(matched) <= index:
        raise SystemExit(f"未找到第 {index + 1} 个 {kind} 候选，实际候选见上方输出")
    return matched[index]


# ===== 表单域 / #占位符 / 浮动文本框 的合成辅助 =====
#
# 这几类是 Word 里真实存在、但 python-docx 没有 API 直接构造的结构，只能拼 XML。
# 形态与真实模板（data/hetong/2026-9-11 免烘烤捣打料合同模版2026.doc）保持一致。

_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def field_xml(name: str, *, bold: bool = False) -> str:
    """一个 FORMTEXT 表单域：骨架 run（begin/instrText/separate/end）+ 承载变量名的结果 run。"""
    rpr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> FORMTEXT </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f"<w:r>{rpr}<w:t>{name}</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def add_field(paragraph, name: str, *, bold: bool = False) -> None:
    """把表单域追加到段落末尾（域是若干并列 run，不是一个 run）。"""
    root = parse_xml(f"<w:root {_NS}>{field_xml(name, bold=bold)}</w:root>")
    for element in list(root):
        paragraph._p.append(element)


def add_textbox(paragraph, inner_paragraph_xml: str) -> None:
    """浮动文本框：mc:Choice（wps）与 mc:Fallback（VML）各存一份内容相同的副本。

    两份都要扫描、都要回写——Word/WPS 读 Choice，LibreOffice（PDF 预览）读 Fallback。
    """
    content = f"<w:txbxContent>{inner_paragraph_xml}</w:txbxContent>"
    paragraph._p.append(
        parse_xml(
            f"<w:r {_NS}><mc:AlternateContent>"
            f'<mc:Choice Requires="wps"><w:drawing><wp:inline><wps:wsp><wps:txbx>'
            f"{content}</wps:txbx></wps:wsp></wp:inline></w:drawing></mc:Choice>"
            f"<mc:Fallback><w:pict><v:shape><v:textbox>{content}</v:textbox></v:shape></w:pict></mc:Fallback>"
            f"</mc:AlternateContent></w:r>"
        )
    )


def build_field_fixture(path: Path) -> None:
    """合成模板：覆盖表单域 / #占位符 / 浮动文本框 / 合并单元格 / 同段多变量 / 同名多落点。"""
    doc = Document()
    doc.add_paragraph("免烘烤捣打料合同模板")

    # 浮动文本框里的表单域（真实模板的"合同编号 / 签订日期"就是这个形态）
    add_textbox(
        doc.add_paragraph(),
        f"<w:p><w:r><w:t>合同编号：</w:t></w:r>{field_xml('BPO_ID')}</w:p>",
    )

    # 表单域 + 中文标签：回写只能替换域本身，「甲方：」必须留下
    add_field(doc.add_paragraph("甲方："), "PARTY_A_NAME", bold=True)
    # 同名变量第二次出现 → 同一个字段、两处落点
    add_field(doc.add_paragraph("单位名称："), "PARTY_A_NAME")

    # 一个段落里两个域（真实模板里最多一段 9 个）
    para = doc.add_paragraph("7.1乙方应提交履约保证金")
    add_field(para, "AMOUNT")
    para.add_run("元，采用下列第")
    add_field(para, "METHOD")
    para.add_run("种方式提交：")

    # #...# 占位符
    doc.add_paragraph("致#供应商名称#：")

    # 合并单元格（跨两列）里的域 + 列头给出主体
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "甲方（签章）"
    table.cell(0, 1).text = "乙方（签章）"
    # 两列的标签文字一模一样：主体只能从列头看出来，正是真实模板「联系地址」的形态
    for column, name in ((0, "A_ADDR"), (1, "B_ADDR")):
        cell_para = table.cell(1, column).paragraphs[0]
        cell_para.add_run("联系地址：")
        add_field(cell_para, name)
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.paragraphs[0].add_run("备注：")
    add_field(merged.paragraphs[0], "MERGED_TAX")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> int:
    fixture = WORK / "fixture.docx"
    standardized = WORK / "standardized.docx"
    build_fixture(fixture)

    print("1) 扫描候选")
    scan = template_scan.scan_document(fixture)
    print("--- 文档大纲 ---")
    print(scan.outline)
    print("--- 候选 ---")
    for item in scan.candidates:
        print(f"  [{item['id']}] {item['kind']:<12} {item['anchor']:<5} {item['context'][:40]!r}")

    kinds = {item["kind"] for item in scan.candidates}
    for kind in ("blank", "label_empty", "jinja", "empty_cell"):
        check(f"识别出 {kind} 候选", kind in kinds, f"实际 kinds={sorted(kinds)}")

    print("\n2) 回写标准化模板")
    fields = [
        {"key": "party_a_name", "label": "甲方名称", "type": "text", "required": True,
         "candidate_id": pick(scan.candidates, "blank", 0)["id"]},
        {"key": "party_b_name", "label": "乙方名称", "type": "text", "required": True,
         "candidate_id": pick(scan.candidates, "blank", 1)["id"]},
        {"key": "contract_amount", "label": "合同金额", "type": "number", "required": True,
         "candidate_id": pick(scan.candidates, "label_empty")["id"]},
        # 已有占位符：走"改名"路径，键由 AI 重新给出语义名
        {"key": "sign_date", "label": "签订日期", "type": "date", "required": True,
         "candidate_id": pick(scan.candidates, "jinja")["id"]},
        {"key": "unit_price", "label": "单价", "type": "number", "required": False,
         "candidate_id": pick(scan.candidates, "empty_cell")["id"]},
    ]
    result = template_rewrite.rewrite_document(fixture, fields, standardized)
    print(f"  写入字段：{result['applied_keys']}")
    print(f"  写入方式：{result['write_modes']}")
    if result["warnings"]:
        print(f"  警告：{result['warnings']}")

    expected = {field["key"] for field in fields}
    check("自检：字段全部生成占位符", set(result["applied_keys"]) == expected)
    check("自检：无遗漏变量", not result["missing_in_document"], str(result["missing_in_document"]))

    print("\n3) 校验产物")
    found = set(word_generator.extract_template_variables(standardized))
    check("docxtpl 解析出的变量与字段一致", found == expected, f"found={sorted(found)}")

    doc = Document(str(standardized))
    paragraphs = [p.text for p in doc.paragraphs]
    check("下划线填空已变为占位符", any("{{ party_a_name }}" in t for t in paragraphs),
          str(paragraphs[:4]))
    check("标签后插入占位符", any("合同金额：{{ contract_amount }}" in t for t in paragraphs))
    check("已有占位符被改名", any("{{ sign_date }}" in t for t in paragraphs))
    cell_text = doc.tables[0].cell(1, 1).text
    check("表格空单元格已填充", "{{ unit_price }}" in cell_text, repr(cell_text))

    # 格式保留：占位符所在 run 必须仍带下划线（原填空是下划线）
    underline_kept = False
    for para in doc.paragraphs:
        for run in para.runs:
            if "{{ party_a_name }}" in run.text and run.font.underline:
                underline_kept = True
    check("占位符继承了下划线格式", underline_kept)

    print("\n4) 用 docxtpl 渲染填值")
    rendered = WORK / "rendered.docx"
    word_generator.render_template(
        str(standardized),
        {"party_a_name": "某钢股份有限公司", "party_b_name": "某某钢铁贸易有限公司",
         "contract_amount": "2350000", "sign_date": "2026-09-14", "unit_price": "4700"},
        str(rendered),
    )
    text = "\n".join(p.text for p in Document(str(rendered)).paragraphs)
    check("渲染结果包含填入的甲方", "某钢股份有限公司" in text)
    check("渲染结果不再是占位符", "{{" not in text)

    print("\n5) 端到端：标准化模板转 PDF 并做中文自检")
    try:
        pdf = asyncio.run(docx_convert.convert_to_pdf(rendered))
        spans, ratio = docx_convert._cjk_ink_ratio(pdf)
        check("生成 PDF", pdf.exists() and pdf.stat().st_size > 0)
        check("中文墨迹正常", ratio >= 0.005, f"ratio={ratio:.2%} spans={spans}")
    except BizError as exc:
        check("生成 PDF", False, exc.message)

    print("\n6) 回归：只给 anchor、不给 candidate_id（应用六的实际输出形态）")
    anchor_only = WORK / "anchor_only.docx"
    result = template_rewrite.rewrite_document(
        fixture,
        [
            {"key": "party_a_name", "label": "甲方名称", "required": True,
             "candidate_id": "", "anchor": pick(scan.candidates, "blank", 0)["anchor"]},
            {"key": "contract_amount", "label": "合同金额", "required": True,
             "candidate_id": "", "anchor": pick(scan.candidates, "label_empty")["anchor"]},
            {"key": "unit_price", "label": "单价", "candidate_id": "",
             "anchor": pick(scan.candidates, "empty_cell")["anchor"]},
        ],
        anchor_only,
    )
    check("按锚点回退可定位到填空处", result["write_modes"]["party_a_name"] == "replace",
          result["write_modes"])
    texts = [p.text for p in Document(str(anchor_only)).paragraphs]
    check("下划线痕迹已清除（占位符未追加在下划线之后）",
          any(t.strip() == "甲方（采购方）：{{ party_a_name }}" for t in texts), str(texts[:4]))
    check("标签后插入位置正确",
          any("合同金额：{{ contract_amount }}" in t for t in texts))

    print("\n7) 异常路径")
    try:
        template_rewrite.rewrite_document(fixture, [], WORK / "x.docx")
        check("空字段清单应报错", False, "未抛异常")
    except BizError as exc:
        check("空字段清单应报错", "为空" in exc.message, exc.message)

    bad = [dict(fields[0], candidate_id="c999")]
    try:
        template_rewrite.rewrite_document(fixture, bad, WORK / "x.docx")
        check("不存在的候选应报错", False, "未抛异常")
    except BizError as exc:
        check("不存在的候选应报错", "锚点" in exc.message, exc.message)

    run_field_cases()

    print()
    if failures:
        print(f"❌ 失败 {len(failures)} 项：{failures}")
        return 1
    print("✅ 全部通过")
    return 0


def run_field_cases() -> None:
    """表单域 / #占位符 / 浮动文本框 / 合并单元格 / 同段多变量 / 同名多落点。"""
    fixture = WORK / "field-fixture.docx"
    standardized = WORK / "field-standardized.docx"
    build_field_fixture(fixture)

    print("\n8) 扫描：表单域 / #占位符 / 文本框 / 合并单元格")
    scan = template_scan.scan_document(fixture)
    by_name = {item["name"]: item for item in scan.candidates if item.get("name")}

    # 域区间精确到变量本身，「甲方：」不算在内
    party = by_name.get("PARTY_A_NAME")
    check("识别到表单域候选", party is not None)
    if party:
        check("域区间只覆盖变量本身", party["text"] == "PARTY_A_NAME"
              and party["context"][party["start"]:party["end"]] == "PARTY_A_NAME",
              f"start={party['start']} end={party['end']} context={party['context']!r}")
    check("识别到 #占位符 候选", by_name.get("供应商名称", {}).get("text") == "#供应商名称#",
          str(by_name.get("供应商名称")))
    check("识别到浮动文本框里的域（外层段落取不到它的文字）",
          "BPO_ID" in by_name)
    # 合并单元格跨两列：去重后应只出一个候选
    merged_hits = [item for item in scan.candidates if item.get("name") == "MERGED_TAX"]
    check("合并单元格只出一个候选", len(merged_hits) == 1, f"实际 {len(merged_hits)} 个")
    bpo = [item for item in scan.candidates if item.get("name") == "BPO_ID"]
    check("文本框两份副本各出一条候选（Choice + Fallback）", len(bpo) == 2,
          f"实际 {len(bpo)} 条")
    check("回退分支标 shadow：候选两条、大纲只出现一次",
          sum(1 for item in bpo if item["shadow"]) == 1 and scan.outline.count("BPO_ID") == 1,
          f"shadow={[item['shadow'] for item in bpo]} outline 出现 {scan.outline.count('BPO_ID')} 次")

    slots = {slot.name: slot for slot in scan.slots if slot.name}
    check("同名变量合并成一个槽位", len(slots["PARTY_A_NAME"].candidate_ids) == 2,
          str(slots["PARTY_A_NAME"].candidate_ids))
    check("主体从列头推断（甲方（签章）/乙方（签章））",
          slots["A_ADDR"].owner_hint == "甲方" and slots["B_ADDR"].owner_hint == "乙方",
          f"{slots['A_ADDR'].owner_hint}/{slots['B_ADDR'].owner_hint}")

    print("\n9) 回写：只替换该替换的，同名写全所有落点")
    # #占位符的内文是中文，key 需要另起英文名；其余槽位沿用原名做 key（与真实链路一致）
    plan = [
        ("PARTY_A_NAME", "PARTY_A_NAME", "甲方单位名称"),
        ("AMOUNT", "AMOUNT", "履约保证金金额"),
        ("METHOD", "METHOD", "提交方式"),
        ("供应商名称", "supplier_name_att", "供应商名称"),
        ("BPO_ID", "BPO_ID", "合同编号"),
        ("A_ADDR", "A_ADDR", "联系地址"),
        ("B_ADDR", "B_ADDR", "联系地址"),
        ("MERGED_TAX", "MERGED_TAX", "税率"),
    ]
    fields = [
        {"key": key, "label": label, "slot_id": slots[name].slot_id,
         "source_name": slots[name].name, "owner": slots[name].owner_hint}
        for name, key, label in plan
    ]

    result = template_rewrite.rewrite_document(fixture, fields, standardized)
    modes = result["write_modes"]
    check("表单域走拆域写入", modes.get("PARTY_A_NAME") == "field", str(modes))
    check("#占位符走整段替换", modes.get("supplier_name_att") == "replace", str(modes))
    check("标签重名已按主体消歧",
          sorted(f["label"] for f in result["fields"] if f["key"] in ("A_ADDR", "B_ADDR"))
          == ["乙方联系地址", "甲方联系地址"],
          str([f["label"] for f in result["fields"]]))

    produced = Document(str(standardized))
    # 段落里的字段有的在正文、有的在表格单元格，断言要两处都收
    texts = [p.text for p in produced.paragraphs]
    texts.extend(
        p.text for table in produced.tables for row in table.rows
        for cell in row.cells for p in cell.paragraphs
    )
    check("中文标签原样保留（甲方：）",
          any(t == "甲方：{{ PARTY_A_NAME }}" for t in texts), str(texts[:6]))
    check("同段两个变量各自落位",
          any(t == "7.1乙方应提交履约保证金{{ AMOUNT }}元，采用下列第{{ METHOD }}种方式提交："
              for t in texts), str(texts))
    check("同名变量写了两处落点",
          sum(t.count("{{ PARTY_A_NAME }}") for t in texts) == 2,
          str([t for t in texts if "PARTY_A_NAME" in t]))
    check("合并单元格里的旁文（备注：）保留",
          any("备注：{{ MERGED_TAX }}" == t for t in texts), str(texts))

    # 文本框两份副本都要写上，否则 LibreOffice 预览会显示旧文本。
    # 注意要读 zip 里的 document.xml —— .docx 是 zip 包，直接解码字节看不到 XML。
    import zipfile  # noqa: PLC0415

    with zipfile.ZipFile(standardized) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    check("表单域骨架已拆干净（无 FORMTEXT / fldChar 残留）",
          "FORMTEXT" not in xml and "fldChar" not in xml)
    check("#占位符已消失", "#供应商名称#" not in xml)
    check("文本框两份副本都写上了占位符", xml.count("{{ BPO_ID }}") == 2,
          f"实际 {xml.count('{{ BPO_ID }}')} 处")
    check("文本框本身没被删掉（绘图元素仍在）", xml.count("txbxContent") == 4,
          f"txbxContent={xml.count('txbxContent')}")

    # 用 docxtpl 渲染一遍，确认真的能填进去
    import docxtpl  # noqa: PLC0415

    document = docxtpl.DocxTemplate(str(standardized))
    document.render({"PARTY_A_NAME": "某钢股份有限公司", "supplier_name_att": "某供应商",
                     "AMOUNT": "10000", "METHOD": "2"})
    rendered = WORK / "field-rendered.docx"
    document.save(str(rendered))
    out = "\n".join(p.text for p in Document(str(rendered)).paragraphs)
    check("渲染结果填入了甲乙方取值", "某钢股份有限公司" in out and "某供应商" in out, out)
    check("渲染结果不再有占位符", "{{" not in out)


if __name__ == "__main__":
    raise SystemExit(main())
